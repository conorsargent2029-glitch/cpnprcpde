"""
chi_swv_to_excel.py

Watches your CHI data folder and fills an Excel workbook one row per run --
one 5-column block (Run, Ch1-4) per frequency in FREQUENCIES (see CONFIG
below), left to right. You're prompted for the workbook's file name each
time you start this script.

How labeling works
-------------------
The CHI software saves two files per run:
  - an auto-named file with the real ip data, e.g. 20260714_172311_SWV.txt
  - a "save:"-named file with your stage label but no readable data,
    e.g. 10Hz_BlankRegen_1.bin (saved at the same moment as the .txt above)

This script matches each .txt to its same-timestamp named .bin to recover the
real stage label (e.g. "BlankRegen"), then reduces it to a base concentration
(blank / 1nM / 10nM -- "Regen"/"time"/cycle-number suffixes are stripped).

A new "round" starts every time the base concentration changes. Within a
round, the first row is labeled with the concentration name and each
following row is labeled in MINS_STEP-minute increments (see CONFIG --
currently 5 mins, 10 mins, ...) -- one row per run, no filler rows for time
points that never happened.

Each cycle's frequencies (e.g. 10/60Hz) are saved back-to-back
within GROUP_GAP_SECONDS of each other, so they're grouped by time
proximity into one row -- NOT by the cycle number in the filename, since
CHI does not guarantee that suffix is chronological. If the macro moves on
before every frequency in FREQUENCIES shows up, the row is written with
whichever frequencies it actually got; the rest are left blank.

Runs with no matching named .bin (stage unknown) default to continuing the
current round, but get highlighted yellow in Excel so you can review them --
occasionally an unlabeled run is genuinely the first reading of a new
concentration rather than a mid-round hiccup, and only you know which.

------------------------------------------------------------------------
ONE-TIME SETUP
------------------------------------------------------------------------
Nothing to edit -- both the data folder and the Excel file are asked for
each time you run this (see DAILY USE below). DEFAULT_WATCH_FOLDER and
DEFAULT_EXCEL_NAME in the CONFIG section are just the values suggested the
very first time, before you've run this and it remembers your last answer.

------------------------------------------------------------------------
DAILY USE
------------------------------------------------------------------------
Double-click "run_watcher.bat" (in the same folder as this script) ONCE at
the start of your session. It'll first ask:
  1. Which folder to pull CHI data from -- press Enter to reuse the last
     one, or paste a new folder path (e.g. a new experiment's save folder).
  2. Which Excel file to write to -- press Enter to reuse the last one, or
     type a new name (e.g. if the old file is open elsewhere, or you want a
     fresh workbook for a new experiment day).
It then backfills any runs already sitting in the folder and keeps watching
for new ones. Leave the window open/minimized. Each completed row prints a
line, e.g.:

    Filled row 5 ("10 mins")

If the Excel file is open in Excel when it tries to save, it'll tell you and
wait -- just close Excel and it saves automatically.

When you're done for the day, just close the window (or Ctrl+C in it).
"""

import re
import sys
import os
import time
import json
import socket
import threading
import http.server
import functools
import webbrowser
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font

# ----------------------- CONFIG -----------------------
EXCEL_DIR = r"C:\Users\gao22\Documents"          # folder the workbook is saved into
DEFAULT_EXCEL_NAME = "SWV_Results.xlsx"          # used if you just press Enter at the prompt
DEFAULT_WATCH_FOLDER = r"C:\Users\gao22\OneDrive\Desktop\Conor\07162026 SWV final figure regen\patch 2"
FREQUENCIES = [10, 60]  # the SWV frequencies your macro runs per cycle, in order --
                        # edit this list if a future protocol uses a different set
MINS_STEP = 5  # minutes added to the label for each run within a round (blank, 5 mins, 10 mins, ...)
POLL_SECONDS = 2.0        # how often to check the folder for new files
LABEL_MATCH_WINDOW = 3.0  # seconds -- how close a named .bin's timestamp must be to the .txt's
GROUP_GAP_SECONDS = 200.0  # max span from the first to the last frequency in one cycle to still
                           # count as the same row (must be comfortably less than the macro's
                           # "delay:" between cycles, or the next cycle's first file gets merged in)
SETTINGS_DIR = r"C:\Users\gao22\Documents"       # where "last used" choices are remembered
DASHBOARD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dashboard")
DASHBOARD_PORT_RANGE = range(8787, 8797)  # tries these in order if one's taken
# --------------------------------------------------------

DASHBOARD_ROWS = []      # in-memory mirror of every row written, for the live dashboard
DASHBOARD_URL = None     # set once the server actually binds a port -- for this machine
DASHBOARD_LAN_URL = None  # same server, reachable from other devices on your network

EXCEL_PATH = None    # set at startup by prompt_for_excel_path()
WATCH_FOLDER = None  # set at startup by prompt_for_watch_folder()


def prompt_for_excel_path():
    """Asks which workbook to write to. Press Enter to use the default name
    (or the file you used last time)."""
    global EXCEL_PATH
    last_used = _load_last_used("chi_last_excel_name")
    default_name = last_used or DEFAULT_EXCEL_NAME
    typed = input(f"Excel file name to write to [{default_name}]: ").strip()
    name = typed or default_name
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    EXCEL_PATH = os.path.join(EXCEL_DIR, name)
    _save_last_used("chi_last_excel_name", name)
    print(f"Writing to: {EXCEL_PATH}\n")


def prompt_for_watch_folder():
    """Asks which folder to pull CHI output files from. Press Enter to use
    the folder you used last time (or the built-in default)."""
    global WATCH_FOLDER
    last_used = _load_last_used("chi_last_watch_folder")
    default_folder = last_used or DEFAULT_WATCH_FOLDER
    while True:
        typed = input(f"Folder to pull CHI data from [{default_folder}]: ").strip().strip('"')
        folder = typed or default_folder
        if os.path.isdir(folder):
            WATCH_FOLDER = folder
            _save_last_used("chi_last_watch_folder", folder)
            print(f"Watching: {WATCH_FOLDER}\n")
            return
        print(f"[!] That folder doesn't exist: {folder}\n")
        default_folder = DEFAULT_WATCH_FOLDER


def _last_used_marker_path(key):
    return os.path.join(SETTINGS_DIR, f".{key}.txt")


def _load_last_used(key):
    try:
        with open(_last_used_marker_path(key), "r") as f:
            return f.read().strip() or None
    except OSError:
        return None


def _save_last_used(key, value):
    try:
        os.makedirs(SETTINGS_DIR, exist_ok=True)
        with open(_last_used_marker_path(key), "w") as f:
            f.write(value)
    except OSError:
        pass

# ------------------------------------------------------------------
# Live browser dashboard (dashboard/dashboard.html polls dashboard/data.json)
# ------------------------------------------------------------------
def get_lan_ip():
    """Best-effort local network IP -- opens a UDP socket to a public
    address without actually sending anything, just to see which local
    interface/IP the OS would route through. Returns None if it can't tell
    (e.g. no network connection)."""
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        return s.getsockname()[0]
    except OSError:
        return None
    finally:
        s.close()


def write_qr_code(url):
    """Writes dashboard/qr.svg for the given URL, so the dashboard page can
    show a scannable code for opening it on your phone. No-op if the
    qrcode package isn't installed -- the dashboard just won't show one."""
    try:
        import qrcode
        import qrcode.image.svg
    except ImportError:
        return False
    img = qrcode.make(url, image_factory=qrcode.image.svg.SvgImage)
    img.save(os.path.join(DASHBOARD_DIR, "qr.svg"))
    return True


def start_dashboard_server():
    """Serves DASHBOARD_DIR so dashboard.html can auto-refresh from
    data.json. Binds to all network interfaces (not just this machine), so
    other devices on your network can reach it too -- sets DASHBOARD_URL
    (this machine) and DASHBOARD_LAN_URL (other devices), or leaves both
    None if no port could be bound."""
    global DASHBOARD_URL, DASHBOARD_LAN_URL
    os.makedirs(DASHBOARD_DIR, exist_ok=True)
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=DASHBOARD_DIR)
    for port in DASHBOARD_PORT_RANGE:
        try:
            server = http.server.ThreadingHTTPServer(("0.0.0.0", port), handler)
        except OSError:
            continue
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        DASHBOARD_URL = f"http://127.0.0.1:{port}/dashboard.html"
        lan_ip = get_lan_ip()
        DASHBOARD_LAN_URL = f"http://{lan_ip}:{port}/dashboard.html" if lan_ip else None
        if DASHBOARD_LAN_URL:
            write_qr_code(DASHBOARD_LAN_URL)
        return DASHBOARD_URL
    print("[!] Couldn't start the live dashboard (no free port found) -- Excel will still update normally.")
    return None


def write_dashboard_data(current_round):
    if DASHBOARD_URL is None:
        return
    data = {
        "updated_at": datetime.now().isoformat(),
        "watch_folder": WATCH_FOLDER,
        "excel_path": EXCEL_PATH,
        "current_round": current_round,
        "frequencies": FREQUENCIES,
        "rows": DASHBOARD_ROWS,
    }
    tmp_path = os.path.join(DASHBOARD_DIR, "data.json.tmp")
    final_path = os.path.join(DASHBOARD_DIR, "data.json")
    try:
        with open(tmp_path, "w") as f:
            json.dump(data, f)
        os.replace(tmp_path, final_path)  # atomic, so the browser never reads a half-written file
    except OSError:
        pass


def seed_dashboard_from_excel():
    """On startup, load whatever's already in the workbook so the dashboard
    shows full history immediately, not just rows written this session."""
    DASHBOARD_ROWS.clear()
    if not os.path.exists(EXCEL_PATH):
        return
    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb["SWV Data"] if "SWV Data" in wb.sheetnames else wb.active
    except OSError:
        return
    max_col = block_start_col(len(FREQUENCIES) - 1) + 4
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=max_col):
        label = row[0].value
        if label in (None, ""):
            continue
        flagged = row[0].fill is not None and row[0].fill.start_color.rgb not in (None, "00000000")
        series = {}
        for i, freq in enumerate(FREQUENCIES):
            col0 = i * 5  # 0-indexed offset into `row`; label is col0, ip cells are col0+1..col0+4
            cells = row[col0 + 1: col0 + 5]
            series[str(freq)] = {
                "ip": [c.value for c in cells],
                "estimated": [bool(c.font and c.font.italic) for c in cells],
            }
        DASHBOARD_ROWS.append({"label": label, "flagged": bool(flagged), "series": series})


FREQ_RE = re.compile(r"Frequency \(Hz\)\s*=\s*([\d.]+)")
CHANNEL_RE = re.compile(r"Channel\s*(\d+):")
IP_RE = re.compile(r"ip\s*=\s*(-?[\d.eE+-]+)A")
DATA_TABLE_HEADER_RE = re.compile(r"^Potential/V.*$", re.MULTILINE)

AUTO_NAME_RE = re.compile(r"^\d{8}_\d{6}_SWV\.bin$", re.IGNORECASE)
HZ_PREFIX_RE = re.compile(r"^\s*\d+\s*Hz_?\s*", re.IGNORECASE)
TRAILING_CYCLE_RE = re.compile(r"_\d+$")

CONC_PATTERNS = [
    (re.compile(r"^10\s*nm", re.IGNORECASE), "10nM"),
    (re.compile(r"^1\s*nm", re.IGNORECASE), "1nM"),
    (re.compile(r"^blank", re.IGNORECASE), "blank"),
]

NUM_FMT = "0.00E+00"
UNLABELED_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

PROCESSED_LOG_NAME = ".chi_processed_files.log"
LOCK_FILE_NAME = ".chi_watcher.lock"

_lock_handle = None  # kept open for the process's lifetime -- see acquire_watch_lock


# ------------------------------------------------------------------
# Single-instance lock -- two watchers on the same folder race and corrupt
# both the round-tracking state and the Excel file (each has its own
# in-memory RoundTracker and processed-file set, so they duplicate rows and
# mislabel them). Only one instance may watch a given folder at a time.
#
# This uses a real OS-level exclusive file lock (msvcrt.locking), not a
# "check if a file exists, then write it" approach -- that pattern has a
# race: two processes can both check at the same instant, both see no lock,
# and both proceed. An OS lock is atomic and is released automatically by
# Windows when the process exits for any reason (including a crash or a
# forced kill), so there's no staleness window to tune either.
# ------------------------------------------------------------------
def _lock_path(folder):
    return os.path.join(folder, LOCK_FILE_NAME)


def acquire_watch_lock(folder):
    global _lock_handle
    import msvcrt
    path = _lock_path(folder)
    try:
        f = open(path, "a+b")
    except OSError as e:
        print(f"[!] Couldn't open lock file {path}: {e}")
        return False
    try:
        # msvcrt.locking() locks starting at the file's CURRENT position, not
        # byte 0 -- and "a+b" mode starts positioned at end-of-file, which
        # differs between processes once one has written a PID. Without this
        # seek, two processes can end up locking different byte ranges and
        # never actually conflict.
        f.seek(0)
        msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
    except OSError:
        try:
            f.close()
        except OSError:
            pass
        print("[!] Another instance is already watching this folder -- close that window first.")
        return False
    try:
        f.seek(0)
        f.truncate()
        f.write(str(os.getpid()).encode())
        f.flush()
    except OSError:
        pass  # non-critical -- the lock itself is what matters, not the PID content
    _lock_handle = f
    return True


def release_watch_lock(folder):
    global _lock_handle
    import msvcrt
    if _lock_handle is None:
        return
    try:
        _lock_handle.seek(0)
        msvcrt.locking(_lock_handle.fileno(), msvcrt.LK_UNLCK, 1)
    except OSError:
        pass
    try:
        _lock_handle.close()
    except OSError:
        pass
    _lock_handle = None


# ------------------------------------------------------------------
# Parsing SWV data
# ------------------------------------------------------------------
def compute_difference_peak_from_raw(text, channel):
    """Fallback for when CHI's own peak-picker didn't report a Difference ip
    for this channel (but Forward/Reverse peaks were found). The raw
    per-point difference-current curve (the "i{ch}d" column in the data
    table) is usually still there -- scan it for the point with the
    largest magnitude, the same thing CHI's own peak-picker would report
    if it hadn't given up. Returns None if the table or column is missing."""
    header_match = DATA_TABLE_HEADER_RE.search(text)
    if not header_match:
        return None
    col_index = 1 + (channel - 1) * 3  # columns: Potential, i1d,i1f,i1r, i2d,i2f,i2r, ...
    best = None
    for line in text[header_match.end():].splitlines():
        line = line.strip()
        if not line or "," not in line:
            continue
        parts = line.split(",")
        if len(parts) <= col_index:
            continue
        try:
            value = float(parts[col_index])
        except ValueError:
            continue
        if best is None or abs(value) > abs(best):
            best = value
    return best


def parse_swv_file(filepath):
    """Returns (frequency_hz, {channel_num: ip_value}, {estimated channel
    numbers}) or (None, {}, set()) if this doesn't look like a CHI SWV
    output file. A channel is in the "estimated" set when CHI's own
    Difference peak was missing and the value came from scanning the raw
    difference curve instead."""
    with open(filepath, "r", errors="ignore") as f:
        text = f.read()

    if "Square Wave Voltammetry" not in text:
        return None, {}, set()

    freq_match = FREQ_RE.search(text)
    frequency = float(freq_match.group(1)) if freq_match else None

    ip_values = {}
    parts = CHANNEL_RE.split(text)
    for i in range(1, len(parts), 2):
        channel_num = int(parts[i])
        block = parts[i + 1]
        diff_idx = block.find("Difference:")
        if diff_idx == -1:
            continue
        fwd_idx = block.find("Forward:", diff_idx)
        diff_block = block[diff_idx: fwd_idx if fwd_idx != -1 else None]
        ip_match = IP_RE.search(diff_block)
        if ip_match:
            ip_values[channel_num] = float(ip_match.group(1))

    estimated = set()
    for ch in (1, 2, 3, 4):
        if ch not in ip_values:
            fallback = compute_difference_peak_from_raw(text, ch)
            if fallback is not None:
                ip_values[ch] = fallback
                estimated.add(ch)

    return frequency, ip_values, estimated


# ------------------------------------------------------------------
# Recovering the real stage label from the paired .bin
# ------------------------------------------------------------------
def clean_stage_name(bin_filename):
    """'10Hz_BlankRegen2_3.bin' -> 'BlankRegen2'. Tolerates an accidental
    double prefix like '60Hz_  10Hz_10nMregen_1.bin' -> '10nMregen'."""
    name = bin_filename
    if name.lower().endswith(".bin"):
        name = name[:-4]
    while True:
        stripped = HZ_PREFIX_RE.sub("", name, count=1)
        if stripped == name:
            break
        name = stripped
    name = TRAILING_CYCLE_RE.sub("", name.strip())
    return name.strip()


def find_label_bin(txt_path, folder):
    """Find the named .bin saved at essentially the same moment as this .txt.
    Uses creation time, not modification time: OneDrive (or CHI re-using a
    "save:" name for a later run) can rewrite a named .bin's *content* long
    after it was first created, which updates its mtime but not when it was
    first saved -- comparing mtimes would then match the .txt to a bin from
    a completely different, much later run. Creation time stays fixed."""
    txt_ctime = os.path.getctime(txt_path)
    best_name, best_diff = None, None
    try:
        entries = os.listdir(folder)
    except OSError:
        return None
    for f in entries:
        if not f.lower().endswith(".bin"):
            continue
        if AUTO_NAME_RE.match(f):
            continue  # this is the auto-named twin of the .txt, not a real label
        fpath = os.path.join(folder, f)
        try:
            diff = abs(os.path.getctime(fpath) - txt_ctime)
        except OSError:
            continue
        if diff <= LABEL_MATCH_WINDOW and (best_diff is None or diff < best_diff):
            best_name, best_diff = f, diff
    return best_name


def get_stage_for_txt(txt_path, folder):
    """Returns the cleaned stage name, or None if no named .bin was found."""
    bin_name = find_label_bin(txt_path, folder)
    if bin_name is None:
        return None
    return clean_stage_name(bin_name)


def base_concentration(stage):
    for pattern, canon in CONC_PATTERNS:
        if pattern.match(stage):
            return canon
    return re.sub(r"\d+$", "", stage).strip() or stage


# ------------------------------------------------------------------
# Excel template -- one 5-column block (Run, Ch1-4) per frequency in
# FREQUENCIES, laid out left to right in that order.
# ------------------------------------------------------------------
def block_start_col(freq_index):
    """1-indexed column where the freq_index'th frequency's block starts."""
    return 1 + freq_index * 5


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "SWV Data"

    for i, freq in enumerate(FREQUENCIES):
        col = block_start_col(i)
        ws.cell(row=1, column=col, value=f"{freq} Hz")
        ws.cell(row=2, column=col, value="Run")
        ws.cell(row=2, column=col + 1, value="Ch1")
        ws.cell(row=2, column=col + 2, value="Ch2")
        ws.cell(row=2, column=col + 3, value="Ch3")
        ws.cell(row=2, column=col + 4, value="Ch4")

    return wb


def find_next_empty_row(ws):
    row = 3
    while ws.cell(row=row, column=1).value not in (None, ""):
        row += 1
    return row


ESTIMATED_FONT = Font(italic=True)


def write_ip_row(ws, row, col_start, ip_values, estimated_channels):
    ip_values = ip_values or {}
    estimated_channels = estimated_channels or set()
    for offset, ch in enumerate((1, 2, 3, 4)):
        cell = ws.cell(row=row, column=col_start + offset, value=ip_values.get(ch))
        cell.number_format = NUM_FMT
        if ch in estimated_channels:
            cell.font = ESTIMATED_FONT


def fill_next_row(label, freq_data, freq_names, flag_for_review, current_round=None):
    """freq_data: {freq: (ip_values_dict, estimated_channels_set)} -- a
    frequency missing from this dict (or with None ip_values) is left blank.
    freq_names: {freq: source_filename} for the console summary line."""
    if os.path.exists(EXCEL_PATH):
        wb = load_workbook(EXCEL_PATH)
        ws = wb["SWV Data"] if "SWV Data" in wb.sheetnames else wb.active
    else:
        wb = build_template()
        ws = wb["SWV Data"]

    row = find_next_empty_row(ws)
    label_cells = []
    for i, freq in enumerate(FREQUENCIES):
        col = block_start_col(i)
        label_cells.append(ws.cell(row=row, column=col, value=label))
        ip_values, estimated_channels = freq_data.get(freq, (None, None))
        write_ip_row(ws, row, col_start=col + 1, ip_values=ip_values, estimated_channels=estimated_channels)
    if flag_for_review:
        for cell in label_cells:
            cell.fill = UNLABELED_FILL

    warned = False
    while True:
        try:
            wb.save(EXCEL_PATH)
            break
        except PermissionError:
            if not warned:
                print(f"[!] Couldn't save {EXCEL_PATH} -- it's open in Excel right now. "
                      f"Close it there; this row will save automatically as soon as you do.")
                warned = True
            time.sleep(3.0)

    flag_note = "  [FLAGGED for review -- no named save found]" if flag_for_review else ""
    est_parts = []
    for freq in FREQUENCIES:
        _, estimated_channels = freq_data.get(freq, (None, None))
        if estimated_channels:
            est_parts.append(f"{freq}Hz Ch" + ",".join(str(c) for c in sorted(estimated_channels)))
    est_note = f"  [ESTIMATED from raw curve: {'; '.join(est_parts)}]" if est_parts else ""
    names_summary = " | ".join(f"{freq}Hz: {freq_names.get(freq) or '(none)'}" for freq in FREQUENCIES)
    print(f'Filled row {row} ("{label}") -- {names_summary}{flag_note}{est_note}')

    DASHBOARD_ROWS.append({
        "label": label,
        "flagged": bool(flag_for_review),
        "series": {
            str(freq): {
                "ip": [(freq_data.get(freq, (None, None))[0] or {}).get(ch) for ch in (1, 2, 3, 4)],
                "estimated": [ch in (freq_data.get(freq, (None, None))[1] or set()) for ch in (1, 2, 3, 4)],
            }
            for freq in FREQUENCIES
        },
    })
    write_dashboard_data(current_round)


# ------------------------------------------------------------------
# Round tracking (turns a stream of (stage, ip10, ip60, name10, name60)
# entries into labeled Excel rows)
# ------------------------------------------------------------------
class RoundTracker:
    def __init__(self):
        self.current_base = None
        self.counter = 0

    def label_for(self, stage):
        """stage is None for a run with no matching named .bin."""
        if stage is not None:
            base = base_concentration(stage)
            if base != self.current_base:
                self.current_base = base
                self.counter = 0
                return base, False
            self.counter += 1
            return f"{self.counter * MINS_STEP} mins", False

        # No label available.
        if self.current_base is None:
            return None, True  # caller must supply a fallback label
        self.counter += 1
        return f"{self.counter * MINS_STEP} mins", True


MINS_LABEL_RE = re.compile(r"^(\d+) mins$")


def infer_round_state(rows):
    """Reconstructs (current_base, counter) from a list of already-written
    dashboard rows, so resuming an existing workbook continues its last
    round instead of starting a new one."""
    base, counter = None, 0
    for r in rows:
        label = r["label"]
        if label.startswith("?"):
            continue  # a pre-round unlabeled row -- doesn't establish a base
        m = MINS_LABEL_RE.match(label)
        if m:
            counter = int(m.group(1)) // MINS_STEP
        else:
            base, counter = label, 0
    return base, counter


def emit_group(tracker, stage, freq_data, freq_names):
    """freq_data: {freq: (ip_values, estimated_channels)} for whichever
    frequencies have shown up so far in this cycle (not necessarily all of
    FREQUENCIES). freq_names: {freq: source_filename}."""
    if not any(ip for ip, _ in freq_data.values()):
        # No real data on any frequency (e.g. a run that got stopped before
        # it produced a result). Discard it silently rather than writing a
        # blank row and burning a spot in the round's minute-counter -- the
        # next real run should just continue (or start a new round) as if
        # this one never happened.
        names_str = ", ".join(f"{f}Hz: {n}" for f, n in freq_names.items())
        print(f"Skipped empty run (no data) -- {names_str or '(no files)'}")
        return
    label, flag = tracker.label_for(stage)
    if label is None:
        any_name = next(iter(freq_names.values()), "unknown")
        label = f"? {any_name}"
    fill_next_row(label, freq_data, freq_names, flag_for_review=flag,
                  current_round=tracker.current_base)


# ------------------------------------------------------------------
# Manual mode
# ------------------------------------------------------------------
def run_manual(file_paths):
    folder = os.path.dirname(os.path.abspath(file_paths[0]))
    freq_data = {}
    freq_names = {}
    stage = None
    for path in file_paths:
        freq, ip, est = parse_swv_file(path)
        if freq is None:
            print(f"Couldn't read frequency from {path} -- skipping this run.")
            return
        freq = int(freq)
        freq_data[freq] = (ip, est)
        freq_names[freq] = os.path.basename(path)
        if stage is None:
            stage = get_stage_for_txt(path, folder)

    tracker = RoundTracker()
    # NOTE: manual mode has no memory of prior rounds -- run this only for
    # one-off reprocessing, not as your primary workflow.
    emit_group(tracker, stage, freq_data, freq_names)


# ------------------------------------------------------------------
# Watch mode
# ------------------------------------------------------------------
def load_processed_log(folder):
    log_path = os.path.join(folder, PROCESSED_LOG_NAME)
    if os.path.exists(log_path):
        with open(log_path, "r") as f:
            return set(line.strip() for line in f if line.strip())
    return set()


def append_processed_log(folder, filename):
    log_path = os.path.join(folder, PROCESSED_LOG_NAME)
    with open(log_path, "a") as f:
        f.write(filename + "\n")


def wait_until_stable(filepath, checks_needed=2, interval=0.5, timeout=30):
    """Waits until a file's size stops changing (i.e. CHI has finished writing it)."""
    last_size = -1
    stable_count = 0
    elapsed = 0.0
    while stable_count < checks_needed and elapsed < timeout:
        try:
            size = os.path.getsize(filepath)
        except OSError:
            size = -1
        if size == last_size and size > 0:
            stable_count += 1
        else:
            stable_count = 0
            last_size = size
        time.sleep(interval)
        elapsed += interval


def run_watcher():
    if not os.path.isdir(WATCH_FOLDER):
        print(f"That folder no longer exists: {WATCH_FOLDER}")
        print("Restart the script and enter a valid folder path.")
        return

    if not acquire_watch_lock(WATCH_FOLDER):
        return

    processed = load_processed_log(WATCH_FOLDER)
    tracker = RoundTracker()
    # Accumulates the frequencies seen so far for the cycle currently in
    # progress: {'freq_data': {freq: (ip, estimated)}, 'freq_names': {freq: name},
    # 'stage': str_or_None, 'last_time': mtime_of_most_recent_file}
    current_group = None

    def flush_group():
        nonlocal current_group
        if current_group is not None:
            emit_group(tracker, current_group["stage"], current_group["freq_data"], current_group["freq_names"])
            current_group = None

    seed_dashboard_from_excel()
    # Pick up round-tracking state (current concentration + minute counter)
    # from whatever's already in the workbook, so a restarted watcher
    # continues the round instead of resetting it.
    tracker.current_base, tracker.counter = infer_round_state(DASHBOARD_ROWS)
    url = start_dashboard_server()
    write_dashboard_data(tracker.current_base)

    print(f"Watching: {WATCH_FOLDER}")
    print(f"Writing results to: {EXCEL_PATH}")
    if url:
        print(f"Live dashboard (this computer): {url}")
        if DASHBOARD_LAN_URL:
            print(f"Live dashboard (phone/other devices on your network): {DASHBOARD_LAN_URL}")
            print("  (If it doesn't load on your phone, Windows Firewall may be blocking it --")
            print("   you may need to allow Python through the firewall for private networks.)")
        try:
            webbrowser.open(url)
        except Exception:
            pass
    print("Leave this window open. Press Ctrl+C to stop.\n")

    try:
        while True:
            try:
                candidates = [f for f in os.listdir(WATCH_FOLDER) if f.lower().endswith(".txt")]
                candidates.sort(key=lambda f: os.path.getmtime(os.path.join(WATCH_FOLDER, f)))
            except FileNotFoundError:
                time.sleep(POLL_SECONDS)
                continue

            for fname in candidates:
                if fname in processed or fname == PROCESSED_LOG_NAME:
                    continue

                fpath = os.path.join(WATCH_FOLDER, fname)
                wait_until_stable(fpath)

                freq, ip, estimated = parse_swv_file(fpath)
                if freq is None:
                    continue  # not a CHI SWV file (or still being written)

                freq = int(freq)
                processed.add(fname)
                append_processed_log(WATCH_FOLDER, fname)

                if freq not in FREQUENCIES:
                    print(f"[!] {fname}: {freq}Hz isn't in FREQUENCIES ({FREQUENCIES}) -- ignored. "
                          f"Edit FREQUENCIES at the top of the script if your protocol changed.")
                    continue

                mtime = os.path.getmtime(fpath)
                stage = get_stage_for_txt(fpath, WATCH_FOLDER)

                if current_group is not None:
                    gap = mtime - current_group["last_time"]
                    if freq in current_group["freq_data"] or gap > GROUP_GAP_SECONDS:
                        # Either this frequency already showed up in the
                        # current cycle (the next cycle has started), or
                        # too much time passed (the macro stopped/moved on)
                        # -- close out whatever the current group has,
                        # complete or not, and start fresh.
                        flush_group()

                if current_group is None:
                    current_group = {"freq_data": {}, "freq_names": {}, "stage": stage, "last_time": mtime}
                else:
                    if current_group["stage"] is None and stage is not None:
                        current_group["stage"] = stage
                    current_group["last_time"] = mtime
                current_group["freq_data"][freq] = (ip, estimated)
                current_group["freq_names"][freq] = fname

            # Flush a group that's been sitting without a new file for a
            # while (the macro stopped, or moved on to something else).
            if current_group is not None and (time.time() - current_group["last_time"]) > GROUP_GAP_SECONDS:
                flush_group()

            time.sleep(POLL_SECONDS)
    except KeyboardInterrupt:
        flush_group()
        print("\nStopped watching.")
    finally:
        release_watch_lock(WATCH_FOLDER)


# ------------------------------------------------------------------
def auto_configure():
    """No-prompt startup for the auto-launch shortcut: reuses whatever
    folder/Excel file were used last time (or the built-in defaults if this
    has never been run before)."""
    global WATCH_FOLDER, EXCEL_PATH
    WATCH_FOLDER = _load_last_used("chi_last_watch_folder") or DEFAULT_WATCH_FOLDER
    excel_name = _load_last_used("chi_last_excel_name") or DEFAULT_EXCEL_NAME
    EXCEL_PATH = os.path.join(EXCEL_DIR, excel_name)
    print(f"[auto-start] Watching: {WATCH_FOLDER}")
    print(f"[auto-start] Writing to: {EXCEL_PATH}")
    if not os.path.isdir(WATCH_FOLDER):
        print(f"[!] That folder doesn't exist -- open run_watcher.bat normally (not the auto-start "
              f"shortcut) once to set a valid folder, then auto-start will remember it next time.")


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] != "--auto":
        prompt_for_excel_path()
        run_manual(sys.argv[1:])
    elif len(sys.argv) == 2 and sys.argv[1] == "--auto":
        auto_configure()
        run_watcher()
    elif len(sys.argv) == 1:
        prompt_for_watch_folder()
        prompt_for_excel_path()
        run_watcher()
    else:
        print("Usage:")
        print("  python chi_swv_to_excel.py                    (watch mode, asks for folder + Excel file)")
        print("  python chi_swv_to_excel.py --auto              (watch mode, no prompts -- reuses last settings)")
        print("  python chi_swv_to_excel.py file1.txt file2.txt [file3.txt ...]  (manual, one cycle's worth of files)")
